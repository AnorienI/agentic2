import sqlite3
import sys
import os
import glob
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from langchain_core.tools import tool

sys.path.append(os.path.join(os.path.dirname(__file__), ".."))
from modules import fetch_fx_a2, fetch_cb_rates_a2, fetch_trade_a2

DB_PATH = "currenflux.db"

def _load_versioned_history(base_name: str) -> pd.DataFrame:
    files = sorted(glob.glob(f"{base_name}_*.csv"))
    return pd.concat([pd.read_csv(f) for f in files], ignore_index=True) if files else pd.DataFrame()

@tool
def fetch_and_load_currenflux_data() -> str:
    """Coleta dados reais de câmbio, taxas de Bancos Centrais e Comex,
    e carrega o histórico completo nas tabelas fx_dashboard, cb_rates e trade_data no SQLite."""
    results = []
    conn = sqlite3.connect(DB_PATH)

    fetch_fx_a2.run()
    fx_df = pd.read_csv("fx_dashboard.csv")
    fx_df.to_sql("fx_dashboard", conn, if_exists="replace", index=False)
    results.append(f"fx_dashboard: {len(fx_df)} linha(s)")

    fetch_cb_rates_a2.run()
    cb_df = _load_versioned_history("cb_rates")
    if not cb_df.empty:
        cb_df.to_sql("cb_rates", conn, if_exists="replace", index=False)
        results.append(f"cb_rates: {len(cb_df)} linha(s) (histórico completo)")

    fetch_trade_a2.run()
    trade_df = _load_versioned_history("trade_data")
    if not trade_df.empty:
        trade_df.to_sql("trade_data", conn, if_exists="replace", index=False)
        results.append(f"trade_data: {len(trade_df)} linha(s) (histórico completo)")

    conn.close()
    return "Dados atualizados: " + " | ".join(results)

@tool
def export_query_to_excel(query_sql: str, file_path: str = "currenflux_report.xlsx") -> str:
    """Executa uma consulta SQL no CurrenFlux e exporta em Excel."""
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query(query_sql, conn)
    conn.close()

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="CurrenFlux", index=False)

    return f"Relatório exportado em: {file_path}"

@tool
def generate_consolidated_spreadsheet(file_path: str = "currenflux_consolidated.xlsx") -> str:
    """USE ESTA FERRAMENTA para gerar o relatório consolidado em Excel. 
    NÃO pergunte o caminho ao usuário; use 'currenflux_consolidated.xlsx' como padrão no diretório atual."""
    conn = sqlite3.connect(DB_PATH)
    
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        # 1. Carrega dados mais recentes de cada fonte
        fx_df = pd.read_sql_query("SELECT * FROM fx_dashboard ORDER BY data DESC LIMIT 1", conn) if os.path.exists("fx_dashboard.csv") else pd.DataFrame()
        cb_df = pd.read_sql_query("SELECT * FROM cb_rates", conn)
        trade_df = pd.read_sql_query("SELECT rowid AS _rid, * FROM trade_data", conn)
        # 2. Coleta das linhas organizadas por grupos
        groups = []

        # Grupo 1: Câmbio
        if not fx_df.empty:
            fx_rows = []
            for col in ["USDBRL", "EURBRL", "GBPBRL", "CNYBRL", "JPYBRL"]:
                if col in fx_df.columns:
                    fx_rows.append({
                        "Categoria": "Taxa de Câmbio",
                        "Indicador / Par": col,
                        "Valor": fx_df[col].iloc[0],
                        "Data / Referência": fx_df["data"].iloc[0] if "data" in fx_df.columns else ""
                    })
            if fx_rows:
                groups.append(fx_rows)

        # Grupo 2: Taxas de Juros dos Bancos Centrais
        if not cb_df.empty:
            cb_rows = []
            latest_cb = cb_df.sort_values("date_updated").groupby("central_bank").last().reset_index()
            for _, row in latest_cb.iterrows():
                cb_rows.append({
                    "Categoria": f"Juros ({row.get('central_bank', '')})",
                    "Indicador / Par": row.get("rate_name", ""),
                    "Valor": f"{row.get('rate_value', '')}%",
                    "Data / Referência": row.get("effective_date", row.get("date_updated", ""))
                })
            if cb_rows:
                groups.append(cb_rows)

        # Grupo 3: Comércio Exterior
        # Grupo 3: Comércio Exterior (formato longo: uma linha por país/fluxo)
        if not trade_df.empty:
            COUNTRY_NAMES = {"BRA": "Brasil", "CHN": "China", "EUR": "Zona do Euro",
                             "IND": "Índia", "USA": "EUA"}
            FLOW_ORDER = ["Exportações (US$ Mi)", "Importações (US$ Mi)", "Saldo Comercial (US$ Mi)"]

            # Pega a linha mais recente de cada país/fluxo (o maior rowid)
            latest = trade_df.sort_values("_rid").groupby(["country", "flow_type"]).last().reset_index()

            trade_rows = []
            for code, name in COUNTRY_NAMES.items():
                for flow in FLOW_ORDER:
                    match = latest[(latest["country"] == code) & (latest["flow_type"] == flow)]
                    if match.empty:
                        continue
                    r = match.iloc[0]
                    trade_rows.append({
                        "Categoria": f"Comex ({name})",
                        "Indicador / Par": flow,
                        "Valor": r["value_usd_mil"],
                        "Data / Referência": r["effective_date"]
                    })
            if trade_rows:
                groups.append(trade_rows)
                
        # 3. Intercala linhas em branco e cabeçalhos entre grupos
        summary_rows = []
        header_row = {"Categoria": "Categoria", "Indicador / Par": "Indicador / Par", "Valor": "Valor", "Data / Referência": "Data / Referência"}
        blank_row = {"Categoria": "", "Indicador / Par": "", "Valor": "", "Data / Referência": ""}

        for i, group in enumerate(groups):
            if i > 0:
                summary_rows.append(blank_row)
                summary_rows.append(header_row)
            summary_rows.extend(group)

        # Escreve a aba principal "Resumo Consolidado"
        if summary_rows:
            summary_df = pd.DataFrame(summary_rows)
            summary_df.to_excel(writer, sheet_name="Resumo Consolidado", index=False)

        # 4. Exporta as abas detalhadas de cada tabela do banco
        tables = [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table';").fetchall()]
        for table_name in tables:
            df = pd.read_sql_query(f"SELECT * FROM {table_name}", conn)
            sheet_title = table_name.replace("_", " ").title()[:31]
            df.to_excel(writer, sheet_name=sheet_title, index=False)

        # 5. Estilização dinâmica com OpenPyXL
        if "Resumo Consolidado" in writer.sheets:
            ws = writer.sheets["Resumo Consolidado"]

            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            regular_font = Font(name="Calibri", size=11)
            thin_border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9")
            )

            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                # Identifica se a linha é um cabeçalho de categoria
                is_header = (row[0].value == "Categoria")
                is_blank = all(cell.value is None or str(cell.value).strip() == "" for cell in row)

                if is_blank:
                    continue

                for cell in row:
                    if is_header:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.font = regular_font
                        cell.border = thin_border
                        if cell.column == 3:  # Valor
                            cell.alignment = Alignment(horizontal="right")
                        elif cell.column == 4:  # Data / Referência
                            cell.alignment = Alignment(horizontal="center")
                        else:
                            cell.alignment = Alignment(horizontal="left")

            # Auto-fit na largura das colunas
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col if cell.value is not None)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    conn.close()
    return f"Relatório consolidado gerado com sucesso em: {file_path}"